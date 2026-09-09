"""Generate sample laboratory result Word files for PULSE case confirmation."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
INDIVIDUAL_DIR = OUT_DIR / "individual"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0xB9, 0x1C, 0x1C)
MUTED = RGBColor(0x47, 0x55, 0x69)
BLACK = RGBColor(0x11, 0x18, 0x27)
GREEN = RGBColor(0x16, 0x65, 0x34)

SAMPLE_PATIENT = {
    "Name": "SANTOS, MARIA L.",
    "Birthday": "03/14/1998",
    "Age": "28",
    "Nationality": "PHILIPPINES",
    "Civil Status": "SINGLE",
    "Gender": "FEMALE",
    "Passport/ID No": "NA",
    "Address": "PUROK MABUHAY, BARANGAY POBLACION",
    "Barangay": "POBLACION",
    "City": "BAGO CITY",
    "Province": "Negros Occidental",
    "Region": "Negros Island Region",
    "Email": "sample.patient@pulse-cho.local",
    "Phone": "+639170000000",
}

STAFF = {
    "performed": [
        ("MARIA C. REYES, RMT", "Medical Technologist, City Health Laboratory"),
        ("JOSE L. SANTOS, RMT", "Medical Technologist, City Health Laboratory"),
    ],
    "verified": ("ANA P. VILLANUEVA, RMT", "Senior Medical Technologist"),
    "noted": ("ROBERTO M. DELGADO, MD, FPSP", "Pathologist / Laboratory Head"),
}

# Diseases confirmed by laboratory testing in PULSE / PIDSR practice.
# Clinical-only conditions (tetanus, AEFI, PSP, ILI as a syndrome) are omitted.
REPORTS = [
    {
        "disease": "Dengue Fever",
        "category": "Category II",
        "title": "DENGUE NS1 Ag & IgG/IgM TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1001",
        "control_number": "LC-2026-DF001",
        "issued": "09/09/2026 10:00",
        "specimen": "Serum",
        "test_type": "Immunochromatographic Rapid Diagnostic Test (RDT) Combo Assay",
        "methodology": (
            "Lateral flow immunoassay for qualitative detection of Dengue NS1 antigen "
            "and IgM/IgG antibodies."
        ),
        "results": [
            ("DENGUE VIRUS NS1 ANTIGEN", "DETECTED"),
            ("DENGUE IgM ANTIBODY", "DETECTED"),
            ("DENGUE IgG ANTIBODY", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR ACUTE DENGUE FEVER INFECTION",
        "remarks": (
            "NS1 reactivity indicates the presence of Dengue virus antigen, confirming "
            "acute infection. IgM positivity confirms recent exposure. IgG negativity "
            "suggests a primary infection. Clinical correlation advised."
        ),
    },
    {
        "disease": "COVID-19",
        "category": "Category I",
        "title": "SARS-CoV-2 RT-PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1002",
        "control_number": "LC-2026-CV002",
        "issued": "09/09/2026 10:05",
        "specimen": "Nasopharyngeal / oropharyngeal swab",
        "test_type": "Real-Time Reverse Transcription Polymerase Chain Reaction (RT-PCR)",
        "methodology": (
            "Nucleic acid amplification targeting SARS-CoV-2 N and ORF1ab genes."
        ),
        "results": [
            ("SARS-CoV-2 RNA (N gene)", "DETECTED"),
            ("SARS-CoV-2 RNA (ORF1ab)", "DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "POSITIVE FOR SARS-CoV-2 (COVID-19) INFECTION",
        "remarks": (
            "Detection of SARS-CoV-2 RNA confirms current COVID-19 infection. "
            "Correlate with onset date, clinical findings, and isolation protocols."
        ),
    },
    {
        "disease": "Measles",
        "category": "Category I",
        "title": "MEASLES IgM ELISA TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1003",
        "control_number": "LC-2026-MS003",
        "issued": "09/09/2026 10:10",
        "specimen": "Serum",
        "test_type": "IgM ELISA",
        "methodology": (
            "Enzyme-linked immunosorbent assay for measles-specific IgM antibodies."
        ),
        "results": [
            ("MEASLES IgM ANTIBODY", "REACTIVE / DETECTED"),
            ("MEASLES IgG ANTIBODY", "NOT DETECTED"),
            ("INDEX VALUE (IgM)", "3.82 (Cut-off 1.10)"),
        ],
        "interpretation": "POSITIVE FOR ACUTE MEASLES INFECTION",
        "remarks": (
            "Measles IgM reactivity in a clinically compatible case confirms acute "
            "measles. Notify PESU/RESU immediately. Clinical correlation advised."
        ),
    },
    {
        "disease": "Cholera",
        "category": "Category II",
        "title": "STOOL CULTURE (VIBRIO CHOLERAE) TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1004",
        "control_number": "LC-2026-CH004",
        "issued": "09/09/2026 10:15",
        "specimen": "Rectal swab / stool in Cary-Blair",
        "test_type": "Culture",
        "methodology": (
            "Bacteriologic culture on TCBS agar with biochemical identification and "
            "serogrouping of Vibrio cholerae."
        ),
        "results": [
            ("VIBRIO CHOLERAE O1", "ISOLATED / DETECTED"),
            ("SEROTYPE", "OGAWA"),
            ("OTHER ENTERIC PATHOGENS", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR CHOLERA (VIBRIO CHOLERAE O1)",
        "remarks": (
            "Isolation of V. cholerae O1 confirms cholera. Initiate case-based "
            "investigation, WASH assessment, and outbreak notification as indicated."
        ),
    },
    {
        "disease": "Typhoid and Paratyphoid Fever",
        "category": "Category II",
        "title": "BLOOD CULTURE & WIDAL TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1005",
        "control_number": "LC-2026-TF005",
        "issued": "09/09/2026 10:20",
        "specimen": "Whole blood / serum",
        "test_type": "Blood Culture",
        "methodology": (
            "Aerobic blood culture with identification of Salmonella enterica. "
            "Widal agglutination performed as a supportive serologic test."
        ),
        "results": [
            ("BLOOD CULTURE", "SALMONELLA TYPHI ISOLATED"),
            ("WIDAL TO (S. TYPHI O)", "1:320"),
            ("WIDAL TH (S. TYPHI H)", "1:160"),
        ],
        "interpretation": "POSITIVE FOR TYPHOID FEVER (SALMONELLA TYPHI)",
        "remarks": (
            "Blood culture isolation is confirmatory. Elevated Widal titers support "
            "recent infection. Clinical correlation and antimicrobial susceptibility "
            "are advised."
        ),
    },
    {
        "disease": "Malaria",
        "category": "Category II",
        "title": "MALARIA RDT & MICROSCOPY TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1006",
        "control_number": "LC-2026-ML006",
        "issued": "09/09/2026 10:25",
        "specimen": "Capillary / venous blood",
        "test_type": "Rapid Diagnostic Test",
        "methodology": (
            "Immunochromatographic RDT for Plasmodium antigens, confirmed by thick "
            "and thin blood film microscopy."
        ),
        "results": [
            ("MALARIA RDT (Pf HRP2)", "DETECTED"),
            ("BLOOD SMEAR (THICK FILM)", "POSITIVE"),
            ("SPECIES (THIN FILM)", "PLASMODIUM FALCIPARUM"),
        ],
        "interpretation": "POSITIVE FOR MALARIA (P. FALCIPARUM)",
        "remarks": (
            "RDT and microscopy concordance confirms falciparum malaria. Initiate "
            "species-specific treatment and case investigation."
        ),
    },
    {
        "disease": "Leptospirosis",
        "category": "Category II",
        "title": "LEPTOSPIRA IgM ELISA TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1007",
        "control_number": "LC-2026-LP007",
        "issued": "09/09/2026 10:30",
        "specimen": "Serum",
        "test_type": "IgM ELISA",
        "methodology": (
            "Enzyme-linked immunosorbent assay for Leptospira-specific IgM antibodies."
        ),
        "results": [
            ("LEPTOSPIRA IgM", "REACTIVE / DETECTED"),
            ("INDEX VALUE", "2.45 (Cut-off 1.00)"),
            ("MAT (IF REFERRED)", "PENDING / NOT DONE"),
        ],
        "interpretation": "POSITIVE FOR ACUTE LEPTOSPIROSIS",
        "remarks": (
            "IgM reactivity in a patient with compatible exposure (floodwater) "
            "supports acute leptospirosis. MAT remains the reference confirmatory "
            "test when available."
        ),
    },
    {
        "disease": "Acute Viral Hepatitis",
        "category": "Category II",
        "title": "VIRAL HEPATITIS SEROLOGY TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1008",
        "control_number": "LC-2026-VH008",
        "issued": "09/09/2026 10:35",
        "specimen": "Serum",
        "test_type": "IgM ELISA",
        "methodology": (
            "Serologic immunoassay panel for hepatitis A IgM, HBsAg, and anti-HCV."
        ),
        "results": [
            ("ANTI-HAV IgM", "REACTIVE / DETECTED"),
            ("HBsAg", "NOT DETECTED"),
            ("ANTI-HCV", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR ACUTE VIRAL HEPATITIS A",
        "remarks": (
            "Anti-HAV IgM confirms acute hepatitis A. Investigate food and water "
            "exposure and household contacts. Clinical correlation advised."
        ),
    },
    {
        "disease": "Diphtheria",
        "category": "Category II",
        "title": "THROAT SWAB CULTURE (C. DIPHTHERIAE) TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1009",
        "control_number": "LC-2026-DP009",
        "issued": "09/09/2026 10:40",
        "specimen": "Throat / nasopharyngeal swab",
        "test_type": "Culture",
        "methodology": (
            "Culture on tellurite medium with identification of Corynebacterium "
            "diphtheriae and toxin testing when available."
        ),
        "results": [
            ("CORYNEBACTERIUM DIPHTHERIAE", "ISOLATED / DETECTED"),
            ("TOXIGENICITY (ELEK / PCR)", "TOXIN-PRODUCING"),
            ("GRAM STAIN", "GRAM-POSITIVE BACILLI SEEN"),
        ],
        "interpretation": "POSITIVE FOR DIPHTHERIA",
        "remarks": (
            "Isolation of toxigenic C. diphtheriae confirms diphtheria. This is an "
            "immediately notifiable event. Start antibiotics, antitoxin assessment, "
            "and close-contact prophylaxis."
        ),
    },
    {
        "disease": "Pertussis",
        "category": "Category II",
        "title": "BORDETELLA PERTUSSIS PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1010",
        "control_number": "LC-2026-PT010",
        "issued": "09/09/2026 10:45",
        "specimen": "Nasopharyngeal swab",
        "test_type": "PCR",
        "methodology": (
            "Real-time PCR for Bordetella pertussis insertion sequences / toxin gene."
        ),
        "results": [
            ("BORDETELLA PERTUSSIS DNA", "DETECTED"),
            ("BORDETELLA PARAPERTUSSIS DNA", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "POSITIVE FOR PERTUSSIS (WHOOPING COUGH)",
        "remarks": (
            "Detection of B. pertussis DNA confirms pertussis. Notify surveillance "
            "and assess household contacts, especially infants."
        ),
    },
    {
        "disease": "Bacterial Meningitis",
        "category": "Category II",
        "title": "CSF CULTURE & GRAM STAIN TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1011",
        "control_number": "LC-2026-BM011",
        "issued": "09/09/2026 10:50",
        "specimen": "Cerebrospinal fluid (CSF)",
        "test_type": "Culture",
        "methodology": (
            "CSF Gram stain, antigen testing, and bacterial culture with identification."
        ),
        "results": [
            ("CSF GRAM STAIN", "GRAM-POSITIVE DIPLOCOCCI SEEN"),
            ("CSF CULTURE", "STREPTOCOCCUS PNEUMONIAE ISOLATED"),
            ("CSF LATEX ANTIGEN (PNEUMOCOCCAL)", "DETECTED"),
        ],
        "interpretation": "POSITIVE FOR BACTERIAL MENINGITIS (S. PNEUMONIAE)",
        "remarks": (
            "Culture confirmation of pneumococcal meningitis. Urgent clinical "
            "management and public-health follow-up are required."
        ),
    },
    {
        "disease": "Meningococcal Disease",
        "category": "Category I",
        "title": "N. MENINGITIDIS PCR / CULTURE TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1012",
        "control_number": "LC-2026-MD012",
        "issued": "09/09/2026 10:55",
        "specimen": "CSF / blood",
        "test_type": "PCR",
        "methodology": (
            "CSF Gram stain, culture, and real-time PCR for Neisseria meningitidis."
        ),
        "results": [
            ("NEISSERIA MENINGITIDIS DNA", "DETECTED"),
            ("CSF / BLOOD CULTURE", "N. MENINGITIDIS ISOLATED"),
            ("SEROGROUP", "SEROGROUP A/B/C (AS REPORTED)"),
        ],
        "interpretation": "POSITIVE FOR MENINGOCOCCAL DISEASE",
        "remarks": (
            "Laboratory confirmation of meningococcal disease. Category I — notify "
            "within 24 hours. Identify close contacts for chemoprophylaxis."
        ),
    },
    {
        "disease": "Anthrax",
        "category": "Category I",
        "title": "BACILLUS ANTHRACIS CULTURE TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1013",
        "control_number": "LC-2026-AN013",
        "issued": "09/09/2026 11:00",
        "specimen": "Skin lesion swab / blood",
        "test_type": "Culture",
        "methodology": (
            "Bacteriologic culture and identification of Bacillus anthracis from "
            "lesion or blood, with referral to a reference laboratory as needed."
        ),
        "results": [
            ("BACILLUS ANTHRACIS", "ISOLATED / DETECTED"),
            ("GRAM STAIN", "GRAM-POSITIVE BOXCAR BACILLI"),
            ("BLOOD CULTURE", "POSITIVE"),
        ],
        "interpretation": "POSITIVE FOR ANTHRAX",
        "remarks": (
            "Isolation of B. anthracis confirms anthrax. Category I event. "
            "Investigate livestock exposure and coordinate with animal health authorities."
        ),
    },
    {
        "disease": "Human Avian Influenza",
        "category": "Category I",
        "title": "INFLUENZA A/H5 RT-PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1014",
        "control_number": "LC-2026-AI014",
        "issued": "09/09/2026 11:05",
        "specimen": "Nasopharyngeal / throat swab",
        "test_type": "PCR",
        "methodology": (
            "Real-time RT-PCR for Influenza A and H5 (avian) hemagglutinin gene, "
            "referred to RITM when indicated."
        ),
        "results": [
            ("INFLUENZA A RNA", "DETECTED"),
            ("INFLUENZA A/H5 RNA", "DETECTED"),
            ("INFLUENZA B RNA", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR HUMAN AVIAN INFLUENZA (A/H5)",
        "remarks": (
            "H5 RNA detection is a Category I event. Isolate the patient, notify "
            "DOH/RESU immediately, and investigate poultry exposure."
        ),
    },
    {
        "disease": "Middle East Respiratory Syndrome",
        "category": "Category I",
        "title": "MERS-CoV RT-PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1015",
        "control_number": "LC-2026-ME015",
        "issued": "09/09/2026 11:10",
        "specimen": "Nasopharyngeal / lower respiratory specimen",
        "test_type": "PCR",
        "methodology": (
            "Real-time RT-PCR targeting MERS-CoV upE and ORF1a genes."
        ),
        "results": [
            ("MERS-CoV RNA (upE)", "DETECTED"),
            ("MERS-CoV RNA (ORF1a)", "DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "POSITIVE FOR MIDDLE EAST RESPIRATORY SYNDROME (MERS-CoV)",
        "remarks": (
            "Confirmed MERS-CoV infection. Category I — immediate notification, "
            "isolation, and contact tracing required."
        ),
    },
    {
        "disease": "Severe Acute Respiratory Syndrome",
        "category": "Category I",
        "title": "SARS-CoV RT-PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1016",
        "control_number": "LC-2026-SA016",
        "issued": "09/09/2026 11:15",
        "specimen": "Nasopharyngeal / lower respiratory specimen",
        "test_type": "PCR",
        "methodology": (
            "Real-time RT-PCR for SARS coronavirus RNA (distinct from SARS-CoV-2)."
        ),
        "results": [
            ("SARS-CoV RNA", "DETECTED"),
            ("SARS-CoV-2 RNA", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "POSITIVE FOR SEVERE ACUTE RESPIRATORY SYNDROME (SARS)",
        "remarks": (
            "Detection of SARS-CoV RNA confirms SARS. Category I emergency "
            "notification and isolation protocols apply."
        ),
    },
    {
        "disease": "Acute Bloody Diarrhea",
        "category": "Category II",
        "title": "STOOL CULTURE (SHIGELLA / DYSENTERY) TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1017",
        "control_number": "LC-2026-BD017",
        "issued": "09/09/2026 11:20",
        "specimen": "Stool / rectal swab",
        "test_type": "Culture",
        "methodology": (
            "Stool culture for Shigella, Salmonella, Campylobacter, and E. coli "
            "with identification of dysentery pathogens."
        ),
        "results": [
            ("SHIGELLA FLEXNERI", "ISOLATED / DETECTED"),
            ("VIBRIO CHOLERAE", "NOT DETECTED"),
            ("FECAL WBC / RBC", "PRESENT"),
        ],
        "interpretation": "POSITIVE FOR ACUTE BLOODY DIARRHEA (SHIGELLOSIS)",
        "remarks": (
            "Isolation of Shigella confirms bacillary dysentery. Investigate "
            "water/food source and household contacts."
        ),
    },
    {
        "disease": "Hand, Foot, and Mouth Disease",
        "category": "Category I",
        "title": "ENTEROVIRUS / COXSACKIE PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1018",
        "control_number": "LC-2026-HF018",
        "issued": "09/09/2026 11:25",
        "specimen": "Throat swab / vesicle fluid / stool",
        "test_type": "PCR",
        "methodology": (
            "RT-PCR for enterovirus and Coxsackievirus A16 / EV71 typing."
        ),
        "results": [
            ("ENTEROVIRUS RNA", "DETECTED"),
            ("COXSACKIEVIRUS A16", "DETECTED"),
            ("EV71", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR HAND, FOOT, AND MOUTH DISEASE (CA16)",
        "remarks": (
            "Enterovirus CA16 detection confirms HFMD. Isolate from school/daycare "
            "until lesions resolve. Watch for neurologic complications."
        ),
    },
    {
        "disease": "Rabies",
        "category": "Category I",
        "title": "RABIES DFA / RT-PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1019",
        "control_number": "LC-2026-RB019",
        "issued": "09/09/2026 11:30",
        "specimen": "Nuchal skin biopsy / saliva / CSF (ante-mortem)",
        "test_type": "PCR",
        "methodology": (
            "Direct fluorescent antibody (DFA) staining and RT-PCR for rabies virus "
            "nucleoprotein gene."
        ),
        "results": [
            ("RABIES VIRUS ANTIGEN (DFA)", "DETECTED"),
            ("RABIES VIRUS RNA (RT-PCR)", "DETECTED"),
            ("NEGATIVE CONTROL", "VALID"),
        ],
        "interpretation": "POSITIVE FOR RABIES",
        "remarks": (
            "Laboratory confirmation of rabies. Category I. Trace animal exposure, "
            "identify other bite victims, and coordinate with animal bite clinic / "
            "veterinary services."
        ),
    },
    {
        "disease": "Acute Flaccid Paralysis",
        "category": "Category I",
        "title": "STOOL VIROLOGY (POLIOVIRUS) TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1020",
        "control_number": "LC-2026-AF020",
        "issued": "09/09/2026 11:35",
        "specimen": "Stool (2 specimens, 24 hours apart)",
        "test_type": "PCR",
        "methodology": (
            "WHO-recommended stool culture / ITD PCR for poliovirus isolation and "
            "intratypic differentiation, referred to the national polio laboratory."
        ),
        "results": [
            ("WILD POLIOVIRUS", "NOT DETECTED"),
            ("VACCINE-DERIVED POLIOVIRUS", "NOT DETECTED"),
            ("NON-POLIO ENTEROVIRUS", "DETECTED"),
        ],
        "interpretation": "LABORATORY-INVESTIGATED ACUTE FLACCID PARALYSIS (NON-POLIO)",
        "remarks": (
            "AFP remains a clinical surveillance diagnosis. Stool virology rules out "
            "wild and VDPV poliovirus. Continue 60-day follow-up as per PIDSR."
        ),
    },
    {
        "disease": "Acute Encephalitis Syndrome",
        "category": "Category II",
        "title": "JAPANESE ENCEPHALITIS IgM / VIRAL PCR TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1021",
        "control_number": "LC-2026-AE021",
        "issued": "09/09/2026 11:40",
        "specimen": "Serum / CSF",
        "test_type": "IgM ELISA",
        "methodology": (
            "JE-specific IgM capture ELISA on serum or CSF, with optional viral PCR panel."
        ),
        "results": [
            ("JAPANESE ENCEPHALITIS IgM (CSF/SERUM)", "REACTIVE / DETECTED"),
            ("DENGUE IgM", "NOT DETECTED"),
            ("CSF BACTERIAL CULTURE", "NO GROWTH"),
        ],
        "interpretation": "POSITIVE FOR ACUTE ENCEPHALITIS SYNDROME (JE IgM REACTIVE)",
        "remarks": (
            "JE IgM in a compatible neurologic case supports Japanese encephalitis "
            "as the AES etiology. Vector-control and immunization history should be reviewed."
        ),
    },
    {
        "disease": "Acute Hemorrhagic Fever Syndrome",
        "category": "Category II",
        "title": "VIRAL HEMORRHAGIC FEVER PANEL TEST REPORT & CERTIFICATION",
        "lab_number": "CHO202609-1022",
        "control_number": "LC-2026-AH022",
        "issued": "09/09/2026 11:45",
        "specimen": "Serum",
        "test_type": "Dengue NS1 Ag & IgG/IgM Combo",
        "methodology": (
            "VHF screening panel: Dengue NS1/IgM/IgG RDT with referral serology as needed."
        ),
        "results": [
            ("DENGUE VIRUS NS1 ANTIGEN", "DETECTED"),
            ("DENGUE IgM ANTIBODY", "DETECTED"),
            ("OTHER VHF PATHOGENS SCREENED", "NOT DETECTED"),
        ],
        "interpretation": "POSITIVE FOR ACUTE HEMORRHAGIC FEVER SYNDROME (DENGUE ETIOLOGY)",
        "remarks": (
            "NS1 and IgM positivity in a hemorrhagic presentation confirms an acute "
            "viral hemorrhagic fever syndrome due to dengue. Clinical correlation advised."
        ),
    },
]

# Matching negative / non-reactive results for the same 22 diseases.
NEGATIVE_DETAILS = {
    "Dengue Fever": {
        "results": [
            ("DENGUE VIRUS NS1 ANTIGEN", "NOT DETECTED"),
            ("DENGUE IgM ANTIBODY", "NOT DETECTED"),
            ("DENGUE IgG ANTIBODY", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR DENGUE FEVER INFECTION",
        "remarks": (
            "NS1, IgM, and IgG were not detected. This result does not support current or "
            "recent dengue infection. If sampled very early after onset, repeat testing may "
            "be considered. Clinical correlation advised."
        ),
    },
    "COVID-19": {
        "results": [
            ("SARS-CoV-2 RNA (N gene)", "NOT DETECTED"),
            ("SARS-CoV-2 RNA (ORF1ab)", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "NEGATIVE FOR SARS-CoV-2 (COVID-19) INFECTION",
        "remarks": (
            "SARS-CoV-2 RNA was not detected. A negative RT-PCR does not completely exclude "
            "infection if the specimen was collected too early or too late. Repeat if clinically indicated."
        ),
    },
    "Measles": {
        "results": [
            ("MEASLES IgM ANTIBODY", "NON-REACTIVE / NOT DETECTED"),
            ("MEASLES IgG ANTIBODY", "NOT DETECTED"),
            ("INDEX VALUE (IgM)", "0.21 (Cut-off 1.10)"),
        ],
        "interpretation": "NEGATIVE FOR ACUTE MEASLES INFECTION",
        "remarks": (
            "Measles IgM is non-reactive. This result does not confirm acute measles. "
            "If rash onset was less than 3 days before collection, a follow-up specimen may be needed."
        ),
    },
    "Cholera": {
        "results": [
            ("VIBRIO CHOLERAE O1", "NOT ISOLATED / NOT DETECTED"),
            ("VIBRIO CHOLERAE O139", "NOT DETECTED"),
            ("OTHER ENTERIC PATHOGENS", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR CHOLERA (VIBRIO CHOLERAE)",
        "remarks": (
            "No Vibrio cholerae was isolated. This result does not confirm cholera. "
            "Culture yield depends on specimen quality and prior antibiotic use."
        ),
    },
    "Typhoid and Paratyphoid Fever": {
        "results": [
            ("BLOOD CULTURE", "NO GROWTH / SALMONELLA NOT ISOLATED"),
            ("WIDAL TO (S. TYPHI O)", "1:20"),
            ("WIDAL TH (S. TYPHI H)", "1:20"),
        ],
        "interpretation": "NEGATIVE FOR TYPHOID AND PARATYPHOID FEVER",
        "remarks": (
            "Blood culture showed no growth of Salmonella Typhi or Paratyphi. Widal titers "
            "are not significant. This result does not confirm enteric fever."
        ),
    },
    "Malaria": {
        "results": [
            ("MALARIA RDT (Pf HRP2)", "NOT DETECTED"),
            ("BLOOD SMEAR (THICK FILM)", "NO MALARIA PARASITE SEEN"),
            ("SPECIES (THIN FILM)", "NONE IDENTIFIED"),
        ],
        "interpretation": "NEGATIVE FOR MALARIA",
        "remarks": (
            "RDT and microscopy are negative for Plasmodium. This result does not confirm malaria. "
            "Repeat smears may be needed if clinical suspicion remains high."
        ),
    },
    "Leptospirosis": {
        "results": [
            ("LEPTOSPIRA IgM", "NON-REACTIVE / NOT DETECTED"),
            ("INDEX VALUE", "0.32 (Cut-off 1.00)"),
            ("MAT (IF REFERRED)", "NOT DONE"),
        ],
        "interpretation": "NEGATIVE FOR ACUTE LEPTOSPIROSIS",
        "remarks": (
            "Leptospira IgM is non-reactive. Early infection may be missed; a convalescent "
            "sample or MAT may be considered if symptoms persist."
        ),
    },
    "Acute Viral Hepatitis": {
        "results": [
            ("ANTI-HAV IgM", "NON-REACTIVE / NOT DETECTED"),
            ("HBsAg", "NOT DETECTED"),
            ("ANTI-HCV", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR ACUTE VIRAL HEPATITIS A/B/C MARKERS",
        "remarks": (
            "No serologic evidence of acute hepatitis A, B, or C on this panel. "
            "Other etiologies of hepatitis should be considered clinically."
        ),
    },
    "Diphtheria": {
        "results": [
            ("CORYNEBACTERIUM DIPHTHERIAE", "NOT ISOLATED / NOT DETECTED"),
            ("TOXIGENICITY (ELEK / PCR)", "NOT APPLICABLE"),
            ("GRAM STAIN", "NO C. DIPHTHERIAE-LIKE BACILLI SEEN"),
        ],
        "interpretation": "NEGATIVE FOR DIPHTHERIA",
        "remarks": (
            "C. diphtheriae was not isolated. This result does not confirm diphtheria. "
            "Prior antibiotics can reduce culture yield; continue clinical assessment if membranes persist."
        ),
    },
    "Pertussis": {
        "results": [
            ("BORDETELLA PERTUSSIS DNA", "NOT DETECTED"),
            ("BORDETELLA PARAPERTUSSIS DNA", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "NEGATIVE FOR PERTUSSIS (WHOOPING COUGH)",
        "remarks": (
            "Bordetella pertussis DNA was not detected. A negative PCR does not fully exclude "
            "pertussis later in the illness. Correlate with clinical and epidemiologic findings."
        ),
    },
    "Bacterial Meningitis": {
        "results": [
            ("CSF GRAM STAIN", "NO ORGANISMS SEEN"),
            ("CSF CULTURE", "NO GROWTH"),
            ("CSF LATEX ANTIGEN (PNEUMOCOCCAL)", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR BACTERIAL MENINGITIS ON THIS SPECIMEN",
        "remarks": (
            "CSF Gram stain, culture, and pneumococcal antigen are negative. This does not "
            "confirm bacterial meningitis. Viral or partially treated bacterial meningitis remains possible."
        ),
    },
    "Meningococcal Disease": {
        "results": [
            ("NEISSERIA MENINGITIDIS DNA", "NOT DETECTED"),
            ("CSF / BLOOD CULTURE", "NO GROWTH"),
            ("SEROGROUP", "NOT APPLICABLE"),
        ],
        "interpretation": "NEGATIVE FOR MENINGOCOCCAL DISEASE",
        "remarks": (
            "N. meningitidis was not detected by PCR or culture. This result does not confirm "
            "meningococcal disease. Early antibiotics can reduce laboratory yield."
        ),
    },
    "Anthrax": {
        "results": [
            ("BACILLUS ANTHRACIS", "NOT ISOLATED / NOT DETECTED"),
            ("GRAM STAIN", "NO BOXCAR BACILLI SEEN"),
            ("BLOOD CULTURE", "NO GROWTH"),
        ],
        "interpretation": "NEGATIVE FOR ANTHRAX",
        "remarks": (
            "Bacillus anthracis was not detected. This result does not confirm anthrax. "
            "Repeat or refer if the epidemiologic link to livestock remains strong."
        ),
    },
    "Human Avian Influenza": {
        "results": [
            ("INFLUENZA A RNA", "NOT DETECTED"),
            ("INFLUENZA A/H5 RNA", "NOT DETECTED"),
            ("INFLUENZA B RNA", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR HUMAN AVIAN INFLUENZA (A/H5)",
        "remarks": (
            "Influenza A/H5 RNA was not detected. A negative PCR does not completely exclude "
            "infection if sampling was suboptimal. Continue public-health assessment if poultry exposure is present."
        ),
    },
    "Middle East Respiratory Syndrome": {
        "results": [
            ("MERS-CoV RNA (upE)", "NOT DETECTED"),
            ("MERS-CoV RNA (ORF1a)", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "NEGATIVE FOR MIDDLE EAST RESPIRATORY SYNDROME (MERS-CoV)",
        "remarks": (
            "MERS-CoV RNA was not detected. This result does not confirm MERS. Repeat lower "
            "respiratory sampling if clinical suspicion remains high."
        ),
    },
    "Severe Acute Respiratory Syndrome": {
        "results": [
            ("SARS-CoV RNA", "NOT DETECTED"),
            ("SARS-CoV-2 RNA", "NOT DETECTED"),
            ("INTERNAL CONTROL", "VALID"),
        ],
        "interpretation": "NEGATIVE FOR SEVERE ACUTE RESPIRATORY SYNDROME (SARS)",
        "remarks": (
            "SARS-CoV RNA was not detected. This result does not confirm SARS. "
            "Correlate with travel, exposure, and other respiratory pathogen tests."
        ),
    },
    "Acute Bloody Diarrhea": {
        "results": [
            ("SHIGELLA SPP.", "NOT ISOLATED / NOT DETECTED"),
            ("VIBRIO CHOLERAE", "NOT DETECTED"),
            ("FECAL WBC / RBC", "NOT SIGNIFICANT"),
        ],
        "interpretation": "NEGATIVE FOR BACTERIAL DYSENTERY PATHOGENS ON THIS CULTURE",
        "remarks": (
            "No Shigella or Vibrio cholerae was isolated. This result does not confirm "
            "acute bloody diarrhea of bacterial etiology. Parasitic or viral causes may still apply."
        ),
    },
    "Hand, Foot, and Mouth Disease": {
        "results": [
            ("ENTEROVIRUS RNA", "NOT DETECTED"),
            ("COXSACKIEVIRUS A16", "NOT DETECTED"),
            ("EV71", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR LABORATORY-CONFIRMED HFMD (CA16 / EV71)",
        "remarks": (
            "Enterovirus, CA16, and EV71 RNA were not detected. HFMD remains a clinical "
            "diagnosis when typical lesions are present; this result does not laboratory-confirm the case."
        ),
    },
    "Rabies": {
        "results": [
            ("RABIES VIRUS ANTIGEN (DFA)", "NOT DETECTED"),
            ("RABIES VIRUS RNA (RT-PCR)", "NOT DETECTED"),
            ("NEGATIVE CONTROL", "VALID"),
        ],
        "interpretation": "NEGATIVE FOR RABIES",
        "remarks": (
            "Rabies antigen and RNA were not detected on this specimen. A negative ante-mortem "
            "test does not fully exclude rabies. Continue clinical management if hydrophobia or bite history is present."
        ),
    },
    "Acute Flaccid Paralysis": {
        "results": [
            ("WILD POLIOVIRUS", "NOT DETECTED"),
            ("VACCINE-DERIVED POLIOVIRUS", "NOT DETECTED"),
            ("NON-POLIO ENTEROVIRUS", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE — NO POLIOVIRUS OR ENTEROVIRUS ISOLATED",
        "remarks": (
            "Stool virology did not isolate wild poliovirus, VDPV, or non-polio enterovirus. "
            "AFP remains a clinical surveillance classification; this laboratory result does not confirm polio."
        ),
    },
    "Acute Encephalitis Syndrome": {
        "results": [
            ("JAPANESE ENCEPHALITIS IgM (CSF/SERUM)", "NON-REACTIVE / NOT DETECTED"),
            ("DENGUE IgM", "NOT DETECTED"),
            ("CSF BACTERIAL CULTURE", "NO GROWTH"),
        ],
        "interpretation": "NEGATIVE FOR JE IgM AND BACTERIAL CSF PATHOGENS",
        "remarks": (
            "No laboratory evidence of Japanese encephalitis or bacterial meningitis on this "
            "panel. AES remains a clinical syndrome; other viral etiologies may still be considered."
        ),
    },
    "Acute Hemorrhagic Fever Syndrome": {
        "results": [
            ("DENGUE VIRUS NS1 ANTIGEN", "NOT DETECTED"),
            ("DENGUE IgM ANTIBODY", "NOT DETECTED"),
            ("OTHER VHF PATHOGENS SCREENED", "NOT DETECTED"),
        ],
        "interpretation": "NEGATIVE FOR DENGUE-ASSOCIATED ACUTE HEMORRHAGIC FEVER",
        "remarks": (
            "Dengue NS1, IgM, and other screened VHF markers were not detected. This result "
            "does not confirm an acute hemorrhagic fever syndrome of dengue etiology."
        ),
    },
}


def as_negative_report(report: dict, index: int) -> dict:
    details = NEGATIVE_DETAILS[report["disease"]]
    negative = dict(report)
    negative.update(details)
    negative["outcome"] = "negative"
    negative["lab_number"] = f"CHO202609-{2001 + index:04d}"
    negative["control_number"] = report["control_number"].replace("LC-2026-", "LC-2026-N")
    negative["issued"] = "09/09/2026 14:00"
    return negative


def set_run_font(run, *, size=11, bold=False, color=BLACK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "4"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "CBD5E1"))
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def add_paragraph(doc, text, *, size=11, bold=False, color=BLACK, align="left", space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_label_value(cell, label, value, *, value_bold=True):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10, bold=True, color=NAVY)
    r2 = p.add_run(str(value))
    set_run_font(r2, size=10, bold=value_bold, color=BLACK)


def build_report_page(doc: Document, report: dict, *, sample_banner=True):
    if sample_banner:
        banner = add_paragraph(
            doc,
            "SAMPLE DOCUMENT — Bago City Health Office / PULSE Case Confirmation Template. "
            "Not an official Red Cross certificate. For surveillance training and OCR testing.",
            size=8,
            bold=True,
            color=RED,
            align="center",
            space_after=6,
        )
        banner.runs[0].italic = True

    add_paragraph(
        doc,
        "BAGO CITY HEALTH OFFICE  ·  CITY HEALTH LABORATORY",
        size=11,
        bold=True,
        color=NAVY,
        align="center",
        space_after=2,
    )
    add_paragraph(
        doc,
        "PULSE Health Surveillance  ·  Negros Occidental, Philippines",
        size=9,
        color=MUTED,
        align="center",
        space_after=8,
    )
    add_paragraph(doc, report["title"], size=14, bold=True, color=NAVY, align="center", space_after=2)
    outcome = report.get("outcome", "positive")
    if outcome == "negative":
        subtitle = f"PIDSR {report['category']}  ·  Laboratory result: NEGATIVE for {report['disease']}"
    else:
        subtitle = f"PIDSR {report['category']}  ·  Confirmed disease: {report['disease']}"
    add_paragraph(doc, subtitle, size=9, color=MUTED, align="center", space_after=10)

    add_paragraph(doc, "PATIENT IDENTIFICATION", size=10, bold=True, color=NAVY, space_after=4)
    id_fields = list(SAMPLE_PATIENT.items())
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [Inches(3.5), Inches(3.5)])
    for i in range(7):
        left = id_fields[i]
        right = id_fields[i + 7]
        add_label_value(table.cell(i, 0), left[0], left[1], value_bold=False)
        add_label_value(table.cell(i, 1), right[0], right[1], value_bold=False)
        for col in range(2):
            set_cell_border(table.cell(i, col), sz="4", color="E2E8F0")

    add_paragraph(doc, "TEST RESULTS", size=10, bold=True, color=NAVY, space_before=12, space_after=4)
    result_table = doc.add_table(rows=len(report["results"]) + 1, cols=2)
    set_table_widths(result_table, [Inches(4.6), Inches(2.4)])
    hdr = result_table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    p0 = hdr[0].paragraphs[0]
    r = p0.add_run("Analyte / Marker")
    set_run_font(r, size=9, bold=True, color=NAVY)
    p1 = hdr[1].paragraphs[0]
    r = p1.add_run("Result")
    set_run_font(r, size=9, bold=True, color=NAVY)
    shade_cell(hdr[0], "E8EEF5")
    shade_cell(hdr[1], "E8EEF5")
    for idx, (analyte, value) in enumerate(report["results"], start=1):
        c0, c1 = result_table.rows[idx].cells
        c0.text = ""
        c1.text = ""
        a = c0.paragraphs[0].add_run(analyte)
        set_run_font(a, size=10, bold=True, color=BLACK)
        positive = any(token in value.upper() for token in ("DETECTED", "REACTIVE", "ISOLATED", "POSITIVE")) and "NOT DETECTED" not in value.upper() and "NO GROWTH" not in value.upper()
        v = c1.paragraphs[0].add_run(value)
        set_run_font(v, size=10, bold=True, color=RED if positive else GREEN)
        for cell in (c0, c1, hdr[0], hdr[1]):
            set_cell_border(cell, sz="4", color="CBD5E1")

    add_paragraph(doc, "INTERPRETATION", size=10, bold=True, color=NAVY, space_before=12, space_after=2)
    interp_color = GREEN if report.get("outcome") == "negative" else RED
    add_paragraph(doc, report["interpretation"], size=12, bold=True, color=interp_color, space_after=8)

    add_paragraph(doc, "REMARKS", size=10, bold=True, color=NAVY, space_after=2)
    add_paragraph(doc, report["remarks"], size=10, color=BLACK, space_after=10)

    add_paragraph(doc, "TECHNICAL DETAILS", size=10, bold=True, color=NAVY, space_after=4)
    tech = doc.add_table(rows=3, cols=1)
    set_table_widths(tech, [Inches(7.0)])
    add_label_value(tech.cell(0, 0), "Test Type", report["test_type"], value_bold=False)
    add_label_value(tech.cell(1, 0), "Methodology", report["methodology"], value_bold=False)
    add_label_value(tech.cell(2, 0), "Specimen", report["specimen"], value_bold=False)
    for i in range(3):
        set_cell_border(tech.cell(i, 0), sz="4", color="E2E8F0")

    note = (
        "Note: This result should be interpreted together with clinical information. "
        "This is a computer-generated sample form for PULSE case confirmation. "
        "No wet signature is required if the document is unaltered. "
        "For verification, contact Bago City Health Office Laboratory."
    )
    add_paragraph(doc, note, size=8, color=MUTED, space_before=8, space_after=12)

    add_paragraph(doc, "PERSONNEL AND VERIFICATION", size=10, bold=True, color=NAVY, space_after=4)
    staff_table = doc.add_table(rows=3, cols=2)
    set_table_widths(staff_table, [Inches(2.2), Inches(4.8)])
    rows = [
        ("Test Performed by", "\n".join(f"{n}\n{t}" for n, t in STAFF["performed"])),
        ("Verified by", f"{STAFF['verified'][0]}\n{STAFF['verified'][1]}"),
        ("Noted by", f"{STAFF['noted'][0]}\n{STAFF['noted'][1]}"),
    ]
    for i, (role, names) in enumerate(rows):
        staff_table.cell(i, 0).text = ""
        staff_table.cell(i, 1).text = ""
        rr = staff_table.cell(i, 0).paragraphs[0].add_run(role)
        set_run_font(rr, size=9, bold=True, color=NAVY)
        nn = staff_table.cell(i, 1).paragraphs[0].add_run(names)
        set_run_font(nn, size=9, bold=False, color=BLACK)
        for col in range(2):
            set_cell_border(staff_table.cell(i, col), sz="4", color="E2E8F0")

    add_paragraph(doc, "ADMINISTRATIVE FOOTER", size=10, bold=True, color=NAVY, space_before=12, space_after=4)
    foot = doc.add_table(rows=1, cols=3)
    set_table_widths(foot, [Inches(2.4), Inches(2.4), Inches(2.2)])
    add_label_value(foot.cell(0, 0), "Lab Number", report["lab_number"])
    add_label_value(foot.cell(0, 1), "Control Number", report["control_number"])
    add_label_value(foot.cell(0, 2), "Certificate Issued", report["issued"])
    for col in range(3):
        set_cell_border(foot.cell(0, col), sz="4", color="E2E8F0")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def slug(name: str) -> str:
    return (
        name.lower()
        .replace("&", "and")
        .replace("/", "-")
        .replace(",", "")
        .replace(" ", "_")
    )


def write_cover(doc: Document, reports, *, outcome="positive"):
    add_paragraph(doc, "BAGO CITY HEALTH OFFICE", size=14, bold=True, color=NAVY, align="center")
    add_paragraph(doc, "PULSE Health Surveillance System", size=12, color=MUTED, align="center")
    heading = (
        "STANDARD LABORATORY RESULT TEMPLATES — NEGATIVE"
        if outcome == "negative"
        else "STANDARD LABORATORY RESULT TEMPLATES — POSITIVE"
    )
    add_paragraph(
        doc,
        heading,
        size=16,
        bold=True,
        color=NAVY,
        align="center",
        space_before=18,
        space_after=8,
    )
    summary = (
        "These forms show NOT DETECTED / non-reactive / no-growth results. "
        "They do not confirm the disease and should not be used to classify a PULSE case as Confirmed."
        if outcome == "negative"
        else "These forms show DETECTED / reactive / isolated results used to verify and confirm the disease."
    )
    add_paragraph(doc, summary, size=11, color=BLACK, align="center", space_after=14)
    add_paragraph(
        doc,
        "Layout matches a Dengue NS1 Ag & IgG/IgM Test Report & Certification: patient identification, "
        "marker results, interpretation, remarks, test type/methodology, personnel, lab number, "
        "control number, and issue date.",
        size=10,
        color=BLACK,
        space_after=10,
    )
    add_paragraph(
        doc,
        "Included diseases" if outcome == "positive" else "Included negative reports",
        size=11,
        bold=True,
        color=NAVY,
        space_after=4,
    )
    for i, report in enumerate(reports, start=1):
        add_paragraph(
            doc,
            f"{i}. {report['disease']} ({report['category']}) — {report['title']}",
            size=10,
            space_after=2,
        )
    add_paragraph(
        doc,
        "Patient demographics are SAMPLE only (Santos, Maria L.).",
        size=9,
        color=MUTED,
        space_before=12,
    )


def save_docx(doc: Document, path: Path) -> Path:
    try:
        doc.save(path)
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        doc.save(alt)
        print(f"Permission denied for {path.name}; wrote {alt.name} instead.")
        return alt


def main():
    missing = [r["disease"] for r in REPORTS if r["disease"] not in NEGATIVE_DETAILS]
    if missing:
        raise SystemExit(f"Missing negative details for: {missing}")

    INDIVIDUAL_DIR.mkdir(parents=True, exist_ok=True)
    negative_dir = INDIVIDUAL_DIR / "negative"
    negative_dir.mkdir(parents=True, exist_ok=True)

    combined = Document()
    configure_document(combined)
    write_cover(combined, REPORTS, outcome="positive")
    for report in REPORTS:
        combined.add_page_break()
        build_report_page(combined, report)
        single = Document()
        configure_document(single)
        build_report_page(single, report)
        save_docx(single, INDIVIDUAL_DIR / f"{slug(report['disease'])}_lab_result.docx")
    combined_path = save_docx(combined, OUT_DIR / "PULSE_Standard_Laboratory_Results_Positive.docx")
    save_docx(combined, OUT_DIR / "PULSE_Standard_Laboratory_Results.docx")

    negatives = [as_negative_report(report, idx) for idx, report in enumerate(REPORTS)]
    combined_neg = Document()
    configure_document(combined_neg)
    write_cover(combined_neg, negatives, outcome="negative")
    for report in negatives:
        combined_neg.add_page_break()
        build_report_page(combined_neg, report)
        single = Document()
        configure_document(single)
        build_report_page(single, report)
        save_docx(single, negative_dir / f"{slug(report['disease'])}_negative_lab_result.docx")
    combined_neg_path = save_docx(
        combined_neg, OUT_DIR / "PULSE_Standard_Laboratory_Results_Negative.docx"
    )

    print(f"Wrote {combined_path}")
    print(f"Wrote {combined_neg_path}")
    print(f"Wrote {len(REPORTS)} positive files in {INDIVIDUAL_DIR}")
    print(f"Wrote {len(negatives)} negative files in {negative_dir}")


if __name__ == "__main__":
    main()
