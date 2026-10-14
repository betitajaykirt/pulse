from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from reports.incident_export import REPORT_HEADERS, REPORT_TITLE, build_incident_workbook


class IncidentExportTests(SimpleTestCase):
    def test_workbook_has_title_and_room_for_full_values(self):
        payload = build_incident_workbook([[
            'Poblacion',
            'Typhoid and Paratyphoid Fever',
            'Confirmed',
            '09/15/2026',
            'Low',
            'Jaykirt Betito',
            '09/17/2026 09:56 AM',
        ]])
        sheet = load_workbook(BytesIO(payload)).active

        self.assertEqual(sheet['A1'].value, REPORT_TITLE)
        self.assertEqual(sheet['G1'].value, None)
        self.assertEqual([sheet.cell(2, column).value for column in range(1, 8)], list(REPORT_HEADERS))
        self.assertEqual(sheet['B3'].value, 'Typhoid and Paratyphoid Fever')
        self.assertEqual(sheet['G3'].value, '09/17/2026 09:56 AM')
        self.assertGreaterEqual(sheet.column_dimensions['B'].width, len('Typhoid and Paratyphoid Fever'))
        self.assertGreaterEqual(sheet.column_dimensions['G'].width, len('09/17/2026 09:56 AM'))
        self.assertGreaterEqual(sheet.column_dimensions['F'].width, len('Jaykirt Betito'))
